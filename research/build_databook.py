"""Build one Excel "databook" of every research input and dataset the twin uses.

Usage: python -m research.build_databook [--out PATH]
Reads only committed repo files (no API calls). Re-run after data or research changes.
"""
import argparse
import json
import re
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
FONT = "맑은 고딕"
HEAD_FILL = PatternFill("solid", fgColor="1F2937")

KO = {
    "quarter": "분기", "trdar_code": "상권코드", "trdar_name": "상권명", "trdar_type_name": "상권유형",
    "service_name": "업종", "sales_amt_total": "매출(분기)", "sales_cnt_total": "결제건수(분기)",
    "sales_amt_weekday": "주중매출", "sales_amt_weekend": "주말매출",
    "sales_amt_t17_21": "17-21시 매출", "sales_amt_t21_24": "21-24시 매출", "sales_amt_t11_14": "11-14시 매출",
    "sales_amt_male": "남성매출", "sales_amt_female": "여성매출",
    "name": "상호", "branch_name": "지점명", "class_code": "소분류코드", "class_name": "소분류",
    "road_addr": "도로명주소", "jibun_addr": "지번주소", "floor": "층", "franchise_flag": "프랜차이즈(지점명 있음)",
    "role": "용도", "lon": "경도", "lat": "위도", "snapshot": "기준일",
    "license_id": "인허가번호", "uptae": "업태", "status": "영업상태", "open_date": "인허가일",
    "close_date": "폐업일", "months_open": "영업개월", "event": "폐업(1)", "closed_within_24m": "24개월내폐업",
    "area_m2": "면적(㎡)", "area_pyeong": "면적(평)", "male_staff": "남성종사자", "female_staff": "여성종사자",
    "criterion_id": "기준ID", "name_ko": "기준(한글)", "name_en": "기준(영문)", "category": "분류",
    "driver_node": "연결된 동인", "unit": "단위", "data_type": "자료형", "method": "방법",
    "source_id": "출처ID", "industry_scope": "업종범위", "in_sim": "시뮬레이션 반영", "notes": "비고",
    "title": "제목", "publisher": "발행처", "year": "연도", "url": "URL", "type": "유형", "reliability": "신뢰도 메모",
    "from_id": "원인", "to_id": "결과", "sign": "방향", "mechanism_ko": "작용 방식", "confidence": "확신도",
}


def csv(path, rename=True, query=None):
    df = pd.read_csv(ROOT / path, dtype=str, keep_default_na=False)
    if query is not None:
        df = df[query(df)]
    for c in df.columns:
        num = pd.to_numeric(df[c], errors="coerce")
        if num.notna().all() and len(df):
            df[c] = num
    return df.rename(columns=KO) if rename else df


def proposed_updates(md_path):
    text = (ROOT / md_path).read_text(encoding="utf-8")
    m = re.search(r"```ya?ml\s*\n(proposed_updates:.*?)```", text, re.S)
    if not m:
        return pd.DataFrame()
    data = yaml.safe_load(m.group(1)).get("proposed_updates") or {}
    rows = []
    for key, v in data.items():
        if isinstance(v, dict):
            rows.append({"항목": key, **{k: (json.dumps(x, ensure_ascii=False) if isinstance(x, (dict, list)) else x)
                                        for k, x in v.items()}})
        else:
            rows.append({"항목": key, "값": v})
    df = pd.DataFrame(rows)
    df.insert(0, "라운드", Path(md_path).stem)
    return df


def flatten(d, prefix=""):
    rows = []
    for k, v in (d or {}).items():
        key = f"{prefix}.{k}" if prefix else str(k)
        if isinstance(v, dict) and not ({"low", "base", "high"} & set(v) or "value" in v):
            rows += flatten(v, key)
        else:
            if isinstance(v, dict):
                rows.append({"파라미터": key, **{kk: (json.dumps(vv, ensure_ascii=False) if isinstance(vv, (dict, list)) else vv)
                                                for kk, vv in v.items()}})
            else:
                rows.append({"파라미터": key, "값": json.dumps(v, ensure_ascii=False) if isinstance(v, list) else v})
    return rows


def hero_sheet():
    wb = load_workbook(ROOT / "research/guesstimate.xlsx", data_only=True)
    ws = wb["HeroStore"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and "criterion_id" in [str(x) for x in r if x])
    df = pd.DataFrame(rows[hdr_i + 1:], columns=[str(c) for c in rows[hdr_i]]).dropna(how="all")
    return df.rename(columns=KO)


def survival():
    d = json.loads((ROOT / "data/processed/survival_wide.json").read_text(encoding="utf-8"))
    rows = []
    for key, label in (("pork_name_only", "돼지고기 구이집 (상호로 판별)"), ("pork_like", "돼지고기 구이집 (상호+상가정보 매칭, 생존자 편향)")):
        v = d.get(key) or {}
        for t in v.get("table", []):
            rows.append({"집단": label, "개월": t.get("month"), "생존율": t.get("survival"), "하한": t.get("ci_low"),
                         "상한": t.get("ci_high"), "관찰 중": t.get("at_risk"), "표본 수": v.get("n"), "폐업 수": v.get("events"),
                         "24개월내 폐업 비율": v.get("share_closed_within_24m"), "면적 중앙값(평)": v.get("median_area_pyeong")})
    for k, v in d.items():
        if isinstance(v, dict) and "table" in v and k not in ("pork_name_only", "pork_like"):
            for t in v["table"]:
                rows.append({"집단": k, "개월": t.get("month"), "생존율": t.get("survival"), "하한": t.get("ci_low"),
                             "상한": t.get("ci_high"), "관찰 중": t.get("at_risk"), "표본 수": v.get("n"), "폐업 수": v.get("events")})
    return pd.DataFrame(rows)


def build(out):
    sheets = [
        ("기준표", "기준(온톨로지) 100개: 브레인스토밍 기준을 매출·이익 동인에 연결", "ontology/criteria.csv",
         csv("ontology/criteria.csv")),
        ("동인 연결", "기준 사이의 원인과 결과, 방향, 근거", "ontology/edges.csv", csv("ontology/edges.csv")),
        ("출처", "모든 숫자의 출처 목록 (논문, 공공데이터, 보고서, 기사)", "ontology/sources.csv", csv("ontology/sources.csv")),
        ("추정값(합성매장)", "합성 매장 값: 낮음/기본/높음, 출처, 확신도, 검토자", "research/guesstimate.xlsx", hero_sheet()),
        ("상권 매출", "서울시 상권분석서비스 추정매출, 한식음식점, 분기별", "data/processed/seoul_sales.csv",
         csv("data/processed/seoul_sales.csv")),
        ("상권 점포", "상권별 한식 점포 수, 개업·폐업", "data/processed/seoul_stores.csv", csv("data/processed/seoul_stores.csv")),
        ("유동인구", "길단위 유동인구, 시간대·연령", "data/processed/seoul_flpop.csv", csv("data/processed/seoul_flpop.csv")),
        ("직장인구", "상권 직장인구", "data/processed/seoul_workers.csv", csv("data/processed/seoul_workers.csv")),
        ("상주인구", "상권 상주인구", "data/processed/seoul_residents.csv", csv("data/processed/seoul_residents.csv")),
        ("집객시설", "대학, 지하철역 등 집객시설", "data/processed/seoul_facilities.csv", csv("data/processed/seoul_facilities.csv")),
        ("고깃집 목록(대학로)", "소상공인 상가정보: 대학로 일대 고기구이·곱창 점포", "data/processed/stores_sbiz.csv",
         csv("data/processed/stores_sbiz.csv")),
        ("인허가·생존(대학로)", "일반음식점 인허가: 인허가일, 폐업일, 면적, 영업개월", "data/processed/licenses.csv",
         csv("data/processed/licenses.csv")),
        ("생존 요약(종로+신촌)", "돼지고기 구이집과 기타 한식 생존 비교", "data/processed/survival_wide.json", survival()),
        ("리서치 2차", "임대료, 돼지고기 도매가, 탄력성, 카드수수료 등 출처 보강", "research/sourcing_round2.md",
         proposed_updates("research/sourcing_round2.md")),
        ("리서치 3차", "원가 1인분, 체류시간, 인원, 인분", "research/sourcing_round3.md",
         proposed_updates("research/sourcing_round3.md")),
        ("리서치 4차", "월별 계절성, 원육 단가, 평판, 점심, 노쇼, 리뷰 이벤트", "research/sourcing_round4_v2.md",
         proposed_updates("research/sourcing_round4_v2.md")),
        ("시즌 파라미터", "게임 시뮬레이션에 쓰는 값과 근거/가정 표시", "research/season_params.yaml",
         pd.DataFrame(flatten(yaml.safe_load((ROOT / "research/season_params.yaml").read_text(encoding="utf-8")))))
        ,
        ("사건 목록", "외부 사건과 밤 결정 상황, 확률, 근거/가정", "research/events.yaml",
         pd.DataFrame(flatten(yaml.safe_load((ROOT / "research/events.yaml").read_text(encoding="utf-8")))))
        ,
    ]
    toc = pd.DataFrame([{"시트": s, "내용": d, "원본 파일": p, "행 수": len(df)} for s, d, p, df in sheets])
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        pd.DataFrame([{"항목": "이 파일", "설명": "대학로 삼겹살집 디지털 트윈이 쓰는 리서치와 데이터 전체 (리포지토리 파일에서 자동 생성)"},
                      {"항목": "주의", "설명": "추정값은 사람 검토 전 예시값, 확신도 L은 출처 없는 가정"},
                      {"항목": "다시 만들기", "설명": "python -m research.build_databook"}]).to_excel(xw, sheet_name="목차", index=False)
        toc.to_excel(xw, sheet_name="목차", index=False, startrow=5)
        for s, _, _, df in sheets:
            df.to_excel(xw, sheet_name=s[:31], index=False)
    wb = load_workbook(out)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name=FONT, size=10, bold=c.font.bold, color=c.font.color)
                c.alignment = Alignment(vertical="top", wrap_text=False)
        hdr_rows = [1, 6] if ws.title == "목차" else [1]
        for r in hdr_rows:
            for c in ws[r]:
                if c.value is not None:
                    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
                    c.fill = HEAD_FILL
        for i, col in enumerate(ws.columns, 1):
            width = max((sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(c.value)) for c in col if c.value is not None), default=6)
            ws.column_dimensions[get_column_letter(i)].width = min(max(8, width * 1.05 + 1), 70)
        ws.freeze_panes = "A7" if ws.title == "목차" else "A2"
    wb.save(out)
    return toc


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT.parent / "share" / "gogi-research-databook.xlsx"))
    a = ap.parse_args()
    t = build(Path(a.out))
    print(t.to_string(index=False))
    print("wrote", a.out)
