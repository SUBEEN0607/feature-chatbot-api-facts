"""Team 1 고깃집 검토표: the shared Google Sheet where members review research values, synced back into the repo.

Run from the repo root (talks to Google only through the `gws` CLI, authenticated as the team lead):

  .venv/bin/python research/sync_review_sheet.py --create    # make the sheet in the Team 1 Drive folder (once)
  .venv/bin/python research/sync_review_sheet.py             # dry run: what the sheet would change (writes nothing)
  .venv/bin/python research/sync_review_sheet.py --apply     # write accepted_patch.yaml, add new sources,
                                                             # rebuild guesstimate.xlsx + hero_base.yaml
  .venv/bin/python research/sync_review_sheet.py --apply --stamp   # also write 반영됨 dates on applied rows
  .venv/bin/python research/sync_review_sheet.py --refresh   # rewrite the 현재 columns from the workbook

Member tabs hold exactly the values docs/TASKS.md gives each member: HeroStore criterion_ids, plus
season_params.yaml / sim/params.py / Benchmarks values (ids prefixed season_params:, sim_params:, BM). A row counts
only when 검토자 is filled and 확인 결과 is 맞음 or 수정:
  맞음  the current value is confirmed: reviewed_by is set on it (a 근거 링크 also sets its source)
  수정  new 낮음/기본/높음 (numbers, low <= base <= high) with a 근거 링크 (S-number or URL) replace it
HeroStore rows go into research/review/accepted_patch.yaml (the shape research/build_workbook.py --apply-patch
reads), merged with the AI-helper rows already there. Other rows go under `sim_assumptions` in that file: reported,
never applied, since they are simulation settings a person edits by hand. The script never writes to the sheet
except the 반영됨 column (--stamp) and the 현재 columns (--refresh).
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import subprocess
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:      # run as a script from anywhere: sim.params must import
    sys.path.insert(0, str(ROOT))
TASKS_MD = ROOT / "docs" / "TASKS.md"
XLSX = ROOT / "research" / "guesstimate.xlsx"
SEASON_PARAMS = ROOT / "research" / "season_params.yaml"
CRITERIA = ROOT / "ontology" / "criteria.csv"
SOURCES = ROOT / "ontology" / "sources.csv"
PATCH = ROOT / "research" / "review" / "accepted_patch.yaml"
SHEET_JSON = ROOT / "research" / "review" / "review_sheet.json"
DATA = ROOT / "data" / "processed"

TITLE = "Team 1 고깃집 검토표"
FOLDER_ID = "1rLSf5zkDB-7E94gBtiogV52duBAzHuSY"     # Drive: Team 1 AI Assignment (its sharing applies)
MEMBERS = ["오하영", "김정연", "박재현", "이지우"]
FONT = "Malgun Gothic"                               # 맑은 고딕
VIA = "검토표"                                        # filled_by of rows this script writes into the patch
HEAD = ["기준ID", "기준", "현재 낮음", "현재 기본", "현재 높음", "단위", "현재 출처", "근거/가정",
        "확인 결과", "수정 낮음", "수정 기본", "수정 높음", "근거 링크", "메모", "검토자", "반영됨"]
CURRENT = slice(1, 8)          # B..H: rewritten by --refresh
INPUT = range(8, 15)           # I..O: members only
STAMP_COL = "P"                # 반영됨: written by --apply --stamp
WIDTHS = [190, 220, 80, 80, 80, 70, 110, 380, 90, 80, 80, 80, 200, 220, 80, 100]
CHOICES = ("맞음", "수정")


# --- gws -------------------------------------------------------------------------------------------------

def gws(*args: str, params: dict | None = None, body: dict | None = None) -> dict:
    """One gws call; JSON in, JSON out. GWS_BIN points tests at a fake."""
    cmd = [os.environ.get("GWS_BIN", "gws"), *args]
    if params is not None:
        cmd += ["--params", json.dumps(params, ensure_ascii=False)]
    if body is not None:
        cmd += ["--json", json.dumps(body, ensure_ascii=False)]
    p = subprocess.run(cmd, capture_output=True, text=True)
    if p.returncode != 0:
        raise RuntimeError(f"gws {' '.join(args)} failed: {p.stderr.strip() or p.stdout.strip()}")
    out = p.stdout.strip()
    return json.loads(out[out.index("{"):]) if "{" in out else {}


def sheet_id() -> str:
    if not SHEET_JSON.exists():
        sys.exit(f"{SHEET_JSON.relative_to(ROOT)} not found: run --create first")
    return json.loads(SHEET_JSON.read_text(encoding="utf-8"))["spreadsheet_id"]


# --- what each member reviews (docs/TASKS.md) ---------------------------------------------------------------

def _hero_rows() -> dict[str, dict]:
    ws = load_workbook(XLSX, read_only=True, data_only=True)["HeroStore"]
    rows = list(ws.iter_rows(values_only=True))
    ix = {h: i for i, h in enumerate(rows[0])}
    return {r[ix["criterion_id"]]: {h: r[i] for h, i in ix.items()} for r in rows[1:] if r[ix["criterion_id"]]}


def _benchmarks() -> dict[str, dict]:
    ws = load_workbook(XLSX, read_only=True, data_only=True)["Benchmarks"]
    rows = list(ws.iter_rows(values_only=True))
    ix = {h: i for i, h in enumerate(rows[0])}
    return {r[ix["bench_id"]]: {h: r[i] for h, i in ix.items()} for r in rows[1:] if r[ix["bench_id"]]}


def _season_leaves(path: str, node, tag=None, note=None):
    """(key, value, tag, note) for every numeric leaf under a season_params.yaml key."""
    if isinstance(node, dict):
        tag, note = node.get("tag", tag), node.get("note", note)
        if "value" in node:
            yield path, node["value"], tag, note
            return
        for k, v in node.items():
            if k in ("tag", "note"):
                continue
            yield from _season_leaves(f"{path}.{k}", v, tag, note)
    elif isinstance(node, (int, float)) and not isinstance(node, bool):
        yield path, node, tag, note


def member_items(tasks_md: Path = TASKS_MD) -> dict[str, list[tuple[str, str]]]:
    """member -> [(row id, group label)] in docs/TASKS.md order. Row ids: criterion_id, season_params:<key>,
    sim_params:<key>, BM<n>. A criterion with no HeroStore row (derived, e.g. drv_covers) is left out."""
    hero = _hero_rows()
    sp = yaml.safe_load(SEASON_PARAMS.read_text(encoding="utf-8"))
    from sim.params import SIM_ASSUMPTIONS
    out: dict[str, list[tuple[str, str]]] = {}
    member = None
    for line in tasks_md.read_text(encoding="utf-8").splitlines():
        m = re.match(r"##\s+(\S+):", line)
        if m:
            member = m.group(1) if m.group(1) in MEMBERS else None
            continue
        if not member or not line.startswith("|") or line.startswith("|---"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) < 2 or cells[0] == "값":
            continue
        label, where = cells[0], cells[1]
        ids: list[str] = []
        for a, b in re.findall(r"`(BM\d+)`\s*부터\s*`(BM\d+)`", where):
            ids += [f"BM{n}" for n in range(int(a[2:]), int(b[2:]) + 1)]
        ctx = None
        for tok in re.findall(r"`([^`]+)`", where):
            if tok.endswith(".yaml") and "season_params" in tok:
                ctx = "season"
            elif tok.endswith(".py") and "params" in tok:
                ctx = "sim"
            elif tok.startswith("BM") or "/" in tok or tok.endswith((".yaml", ".py", ".csv", ".md")):
                continue
            elif tok in hero:
                ids.append(tok)
            elif ctx == "season":
                node = sp
                for part in tok.split("."):
                    node = (node or {}).get(part) if isinstance(node, dict) else None
                ids += [f"season_params:{k}" for k, *_ in _season_leaves(tok, node)]
            elif ctx == "sim" and tok in SIM_ASSUMPTIONS:
                ids.append(f"sim_params:{tok}")
        seen = {i for i, _ in out.get(member, [])}
        out.setdefault(member, []).extend((i, label) for i in dict.fromkeys(ids) if i not in seen)
    return out


def current_values() -> dict[str, dict]:
    """row id -> the 현재 columns (B..H) from the workbook and the sim settings."""
    crit = {r["criterion_id"]: r for r in csv.DictReader(CRITERIA.open(encoding="utf-8"))}
    out = {}
    for cid, h in _hero_rows().items():
        src = h.get("source_id") or ""
        kind = "근거" if src else "가정"
        out[cid] = {"name": h.get("name_ko") or crit.get(cid, {}).get("name_ko", ""), "low": h["low"], "base": h["base"],
                    "high": h["high"], "unit": h.get("unit") or "", "source": src,
                    "why": f"{kind} (확신도 {h.get('confidence') or '-'}): {h.get('reasoning_chain') or ''}".strip(),
                    "confidence": h.get("confidence"), "hero": True}
    for bid, b in _benchmarks().items():
        src = b.get("source_id") or ""
        out[bid] = {"name": b.get("name_ko") or "", "low": b.get("low"), "base": b.get("value"), "high": b.get("high"),
                    "unit": b.get("unit") or "", "source": src,
                    "why": f"{'근거' if src else '가정'}: {' '.join(str(x) for x in (b.get('basis'), b.get('note')) if x)}",
                    "hero": False}
    sp = yaml.safe_load(SEASON_PARAMS.read_text(encoding="utf-8"))
    for top, node in sp.items():
        for key, val, tag, note in _season_leaves(top, node):
            srcs = ";".join(t for t in re.split(r"[;,\s]+", str(tag or "")) if re.fullmatch(r"S\d+", t))
            kind = ("근거+가정" if "가정" in str(tag) else "근거") if srcs else "가정"
            out[f"season_params:{key}"] = {
                "name": key, "low": None, "base": val, "high": None,
                "unit": "₩" if key.endswith(("cost", "_cost", "per_review")) else "", "source": srcs,
                "why": f"{kind}: {note or tag or ''}".strip(), "hero": False}
    from sim.params import SIM_ASSUMPTIONS
    for k, v in SIM_ASSUMPTIONS.items():
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            out[f"sim_params:{k}"] = {"name": k, "low": None, "base": v, "high": None, "unit": "분" if k.endswith("_min") else "",
                                      "source": "", "why": "가정: sim/params.py 의 메커니즘 상수 (출처 없음)", "hero": False}
    return out


def member_rows(member: str, items: dict, cur: dict) -> list[list]:
    rows = []
    for rid, label in items.get(member, []):
        c = cur[rid]
        name = c["name"] if c["hero"] or rid.startswith("BM") else f"{label}: {c['name']}"
        rows.append([rid, name, c["low"], c["base"], c["high"], c["unit"], c["source"], c["why"]])
    return rows


# --- reading the sheet and turning it into a patch ---------------------------------------------------------------

def read_member_tabs(sid: str) -> dict[str, list[dict]]:
    """member -> [{row: sheet row number, 기준ID, ..., 검토자}] keyed by header names (rows found by column A)."""
    res = gws("sheets", "spreadsheets", "values", "batchGet",
              params={"spreadsheetId": sid, "ranges": [f"'{m}'!A1:P500" for m in MEMBERS],
                      "valueRenderOption": "UNFORMATTED_VALUE"})
    out = {}
    for m, vr in zip(MEMBERS, res.get("valueRanges") or []):
        vals = vr.get("values") or []
        head = [str(h).strip() for h in (vals[0] if vals else HEAD)]
        rows = []
        for n, r in enumerate(vals[1:], start=2):
            d = {h: (r[i] if i < len(r) else "") for i, h in enumerate(head)}
            if str(d.get("기준ID") or "").strip():
                d["_row"] = n
                rows.append(d)
        out[m] = rows
    return out


def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v or "").strip().replace(",", "").replace("₩", "").replace("원", "")
    try:
        f = float(s)
    except ValueError:
        return None
    return int(f) if f.is_integer() else f


def _sources() -> list[dict]:
    return list(csv.DictReader(SOURCES.open(encoding="utf-8")))


def resolve_link(link: str, sources: list[dict], new: dict) -> tuple[str | None, str | None]:
    """근거 링크 -> (source_id string, error). S-numbers must exist; a URL reuses a source with that URL or
    becomes a new source (added to ontology/sources.csv on --apply)."""
    link = str(link or "").strip()
    if not link:
        return None, None
    ids = re.findall(r"\bS\d+\b", link)
    urls = re.findall(r"https?://\S+", link)
    known = {s["source_id"] for s in sources}
    bad = [i for i in ids if i not in known]
    if bad:
        return None, f"모르는 출처 번호 {', '.join(bad)}"
    by_url = {s["url"].strip(): s["source_id"] for s in sources if s.get("url")}
    for u in urls:
        if u in by_url:
            ids.append(by_url[u])
        else:
            if u not in new:
                n = max([int(s["source_id"][1:]) for s in sources if re.fullmatch(r"S\d+", s["source_id"])]
                        + [int(v[1:]) for v in new.values()] + [0]) + 1
                new[u] = f"S{n:02d}"
            ids.append(new[u])
    if not ids:
        return None, "근거 링크는 S로 시작하는 출처 번호나 URL"
    return ";".join(dict.fromkeys(ids)), None


def collect(tabs: dict[str, list[dict]], cur: dict, sources: list[dict]) -> dict:
    """Validate member rows. Returns {hero: {cid: patch row}, sim: {id: proposal}, new_sources: {url: id},
    applied: [(member, row, id)], refused: [(member, row, id, reason)], skipped: n}."""
    hero, sim, new_sources, applied, refused, skipped = {}, {}, {}, [], [], 0
    seen: dict[str, tuple] = {}
    for member, rows in tabs.items():
        for r in rows:
            rid = str(r.get("기준ID")).strip()
            who = str(r.get("검토자") or "").strip()
            verdict = str(r.get("확인 결과") or "").strip()
            if not who and not verdict:
                skipped += 1
                continue

            def no(reason):
                refused.append((member, r["_row"], rid, reason))
            if not who:
                no("검토자가 비어 있음")
                continue
            if verdict not in CHOICES:
                no("확인 결과가 맞음 또는 수정이 아님")
                continue
            if rid not in cur:
                no("모르는 기준ID")
                continue
            c = cur[rid]
            link = r.get("근거 링크")
            src, err = resolve_link(link, sources, new_sources)
            if err:
                no(err)
                continue
            memo = str(r.get("메모") or "").strip()
            if verdict == "수정":
                lo, base, hi = (_num(r.get(k)) for k in ("수정 낮음", "수정 기본", "수정 높음"))
                if base is None:
                    no("수정 기본이 숫자가 아님")
                    continue
                if (str(r.get("수정 낮음") or "").strip() and lo is None) or (str(r.get("수정 높음") or "").strip() and hi is None):
                    no("수정 낮음 또는 수정 높음이 숫자가 아님")
                    continue
                if c["unit"] == "범주":       # categorical (e.g. 원산지): text values, no ordering
                    lo, base, hi = (str(r.get(k) or "").strip() or None for k in ("수정 낮음", "수정 기본", "수정 높음"))
                    if not base or not src:
                        no("수정 기본과 근거 링크가 필요함")
                        continue
                    row = {"low": lo or base, "base": base, "high": hi or base, "source_id": src,
                           "reasoning_chain": memo or f"{c['why']} (검토표에서 {who} 수정)"}
                    row.update({"filled_by": VIA, "reviewed_by": who})
                    hero[rid] = row
                    applied.append((member, r["_row"], rid))
                    continue
                if c["hero"] and (lo is None or hi is None):
                    no("수정이면 낮음, 기본, 높음을 모두 숫자로")
                    continue
                if not ((lo is None or lo <= base) and (hi is None or base <= hi)):
                    no("낮음 <= 기본 <= 높음 이 아님")
                    continue
                if not src:
                    no("수정에는 근거 링크가 필요함")
                    continue
                row = {"low": lo, "base": base, "high": hi, "source_id": src,
                       "reasoning_chain": memo or f"{c['why']} (검토표에서 {who} 수정)"}
                if c.get("confidence") == "L":
                    row["confidence"] = "M"
            else:
                row = {}
                if src:
                    row["source_id"] = src if not c["source"] else ";".join(dict.fromkeys(c["source"].split(";") + src.split(";")))
                    if c.get("confidence") == "L":
                        row["confidence"] = "M"
                elif c["hero"] and not c["source"] and c.get("confidence") != "L":
                    no("현재 값에 출처가 없음: 근거 링크를 적어 주세요")
                    continue
            row.update({"filled_by": VIA, "reviewed_by": who})
            key = (verdict, json.dumps({k: row.get(k) for k in ("low", "base", "high", "source_id")}, sort_keys=True))
            if rid in seen and seen[rid][0] != key:     # two tabs disagree on one value: neither is applied
                first = seen[rid][1]
                refused.append((member, r["_row"], rid, f"{first[0]} 탭의 같은 기준과 결과가 다름"))
                refused.append((*first, rid, f"{member} 탭의 같은 기준과 결과가 다름"))
                (hero if c["hero"] else sim).pop(rid, None)
                applied[:] = [a for a in applied if a[2] != rid]
                continue
            seen[rid] = (key, (member, r["_row"]))
            if c["hero"]:
                hero[rid] = row
            else:
                sim[rid] = {"verdict": verdict, **{k: v for k, v in row.items() if k != "filled_by"},
                            "current": c["base"], "note": memo}
            applied.append((member, r["_row"], rid))
    return {"hero": hero, "sim": sim, "new_sources": new_sources, "applied": applied, "refused": refused,
            "skipped": skipped}


def merged_patch(result: dict, existing: dict | None) -> dict:
    """accepted_patch.yaml content: AI-helper rows kept, this sheet's rows replaced by the current sheet."""
    existing = existing or {}
    keep = {k: v for k, v in existing.items()
            if k not in ("_meta", "sim_assumptions") and isinstance(v, dict) and v.get("filled_by") != VIA}
    doc = {"_meta": {
        "generated_by": "research/sync_review_sheet.py (Team 1 고깃집 검토표) + agents.guesstimate",
        "rule": "only rows with 검토자 (reviewed_by) and 확인 결과 맞음 or 수정",
        "note": "Apply with research/build_workbook.py --apply-patch <this file> (sync_review_sheet.py --apply does it)",
        "synced_at": dt.datetime.now().isoformat(timespec="seconds")}}
    doc.update(keep)
    for cid, row in result["hero"].items():
        doc[cid] = row
    if result["sim"]:
        doc["sim_assumptions"] = result["sim"]
    return doc


def report(result: dict) -> str:
    lines = [f"rows: {len(result['applied'])} accepted, {len(result['refused'])} refused, {result['skipped']} not reviewed yet"]
    for m, n, rid, why in result["refused"]:
        lines.append(f"  REFUSED {m} row {n} {rid}: {why}")
    for cid, row in result["hero"].items():
        lines.append(f"  hero {cid}: " + ", ".join(f"{k}={v}" for k, v in row.items() if k != "filled_by"))
    for rid, row in result["sim"].items():
        lines.append(f"  sim setting {rid} (reported, edit by hand): {row['verdict']} {row.get('base', row['current'])}")
    for url, s in result["new_sources"].items():
        lines.append(f"  new source {s}: {url}")
    return "\n".join(lines)


def add_sources(new: dict, who: dict[str, str]) -> None:
    if not new:
        return
    with SOURCES.open(encoding="utf-8") as fh:
        fields = next(csv.reader(fh))
    today = dt.date.today().isoformat()
    with SOURCES.open("a", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        for url, sid in new.items():
            w.writerow({"source_id": sid, "title": "검토표 근거 링크 (제목 보강 필요)", "publisher_author": "",
                        "year": today[:4], "url": url, "type": "web",
                        "reliability_note": f"검토표에서 {who.get(url, '팀원')} 확인 {today}; 제목과 발행처는 다음 검토 때 채움"})


def _reviewed_in_xlsx() -> set[str]:
    return {cid for cid, h in _hero_rows().items() if str(h.get("reviewed_by") or "").strip()}


def _run_build(cmd: list[str], **kw):
    return subprocess.run(cmd, **kw)


def apply(result: dict, run=None) -> None:
    run = run or _run_build
    existing = yaml.safe_load(PATCH.read_text(encoding="utf-8")) if PATCH.exists() else {}
    doc = merged_patch(result, existing)
    lost = _reviewed_in_xlsx() - {k for k in doc if k not in ("_meta", "sim_assumptions")}
    if lost:   # build_workbook rebuilds from its seed values + this patch; reviewed rows outside the patch would vanish
        sys.exit(f"refusing: guesstimate.xlsx has reviewed rows that no patch carries: {sorted(lost)}")
    who = {}
    for cid, row in result["hero"].items():
        for u, s in result["new_sources"].items():
            if s in str(row.get("source_id")):
                who[u] = row["reviewed_by"]
    add_sources(result["new_sources"], who)
    PATCH.parent.mkdir(parents=True, exist_ok=True)
    PATCH.write_text(yaml.safe_dump(doc, allow_unicode=True, sort_keys=False), encoding="utf-8")
    p = run([sys.executable, str(ROOT / "research" / "build_workbook.py"), "--apply-patch", str(PATCH), "--force"],
            cwd=ROOT, capture_output=True, text=True)
    print(p.stdout)
    if p.returncode != 0:
        sys.exit(f"build_workbook failed:\n{p.stderr}")
    print("next: make load (DB), .venv/bin/python web/scripts/sync_data.py --no-sim (검토 N/82 on the web)")


def stamp(sid: str, applied: list[tuple]) -> None:
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    data = [{"range": f"'{m}'!{STAMP_COL}{n}", "values": [[now]]} for m, n, _ in applied]
    if data:
        gws("sheets", "spreadsheets", "values", "batchUpdate", params={"spreadsheetId": sid},
            body={"valueInputOption": "RAW", "data": data})


def refresh(sid: str) -> int:
    """Rewrite B..H (기준 .. 근거/가정) of every member row from the workbook, matched by 기준ID in column A.
    Rows TASKS.md added since --create are appended below the last row (A..H only)."""
    cur, items, tabs = current_values(), member_items(), read_member_tabs(sid)
    data, n = [], 0
    for m in MEMBERS:
        have = {str(r["기준ID"]).strip(): r["_row"] for r in tabs.get(m, [])}
        last = max(have.values(), default=1)
        for row in member_rows(m, items, cur):
            rid = row[0]
            if rid in have:
                data.append({"range": f"'{m}'!B{have[rid]}:H{have[rid]}", "values": [row[CURRENT]]})
            else:
                last += 1
                data.append({"range": f"'{m}'!A{last}:H{last}", "values": [row]})
            n += 1
    gws("sheets", "spreadsheets", "values", "batchUpdate", params={"spreadsheetId": sid},
        body={"valueInputOption": "RAW", "data": data})
    return n


# --- --create ------------------------------------------------------------------------------------------------

GUIDE = [
    ["Team 1 고깃집 검토표"],
    ["대학로 삼겹살집 디지털 트윈에 들어가는 숫자를 팀원이 직접 확인하는 표"],
    [""],
    ["내 탭"],
    ["이름 탭에 내가 맡은 값만 있음, 다른 사람 탭은 고치지 않기"],
    ["회색 칸은 지금 저장소에 들어 있는 값이라 고치지 않기, 흰 칸(확인 결과부터 검토자까지)만 입력"],
    [""],
    ["3단계: 찾기, 확인, 기입"],
    ["1 찾기: 현재 값과 현재 출처를 보고 무료 공개 자료(논문, 정부 통계, 기사), 현장 관찰, 지인 질문으로 숫자 찾기"],
    ["2 확인: 출처를 직접 열어 숫자와 조건(지역, 연도, 단위)이 맞는지 보기"],
    ["3 기입: 확인 결과에서 맞음 또는 수정을 고르고 검토자에 내 이름 쓰기"],
    [""],
    ["맞음과 수정"],
    ["맞음: 현재 값이 맞다고 확인한 경우, 새 출처를 찾았으면 근거 링크에 적기"],
    ["수정: 수정 낮음, 수정 기본, 수정 높음을 모두 숫자로 적기 (낮음 ≤ 기본 ≤ 높음, 단위는 현재 값과 같게)"],
    ["수정에는 근거 링크가 꼭 필요함: S로 시작하는 출처 번호(출처 탭) 또는 자료 URL"],
    ["메모: 계산 근거 한 줄, 수정이면 이 문장이 시뮬레이션의 근거 설명으로 들어감"],
    ["수정 값 칸이 빨갛게 되면 숫자가 아니거나 낮음, 기본, 높음 순서가 맞지 않는 것"],
    [""],
    ["반영 규칙"],
    ["검토자가 비어 있는 행은 반영되지 않음, 검토자와 확인 결과가 모두 있어야 반영"],
    ["김건희가 동기화 스크립트로 저장소에 반영하고, 반영된 행은 반영됨 칸에 날짜가 찍힘"],
    ["기준ID가 season_params, sim_params, BM 으로 시작하는 행은 시뮬레이션 설정 값이라 Claude 가 따로 반영함"],
    ["같은 기준이 두 사람 탭에 있으면 두 사람의 결과가 같아야 반영됨"],
    [""],
    ["지킬 것"],
    ["네이버 플레이스, 캐치테이블, 테이블링은 자동 수집 금지, 직접 보고 소량만 기록"],
    ["가게 이름, 사장님 성함, 전화번호는 적지 않고 T01 같은 번호로만 기록"],
    ["유료 자료와 유료 API 는 쓰지 않음"],
    ["AI 가 찾아 준 숫자는 출처를 직접 열어 확인한 뒤에만 기입"],
    [""],
    ["참고 탭 (읽기 전용)"],
    ["상권 매출: 대학로 상권(3120012) 한식음식점 분기별 추정매출, 서울시 상권분석서비스"],
    ["고깃집 목록: 소상공인 상가정보의 대학로 일대 고기구이와 곱창 점포"],
    ["생존 요약: 종로구와 신촌역 상권의 인허가 기준 생존 분석 요약"],
    ["출처: 저장소의 모든 출처 번호와 링크"],
]


def reference_tabs() -> dict[str, list[list]]:
    tabs = {}
    rows = list(csv.DictReader((DATA / "seoul_sales.csv").open(encoding="utf-8")))
    rows = [r for r in rows if r["trdar_code"] == "3120012"]
    cols = [("quarter", "분기"), ("service_name", "업종"), ("sales_amt_total", "매출(분기, 원)"),
            ("sales_cnt_total", "결제건수"), ("sales_amt_weekday", "주중 매출"), ("sales_amt_weekend", "주말 매출"),
            ("sales_amt_t11_14", "11-14시"), ("sales_amt_t17_21", "17-21시"), ("sales_amt_t21_24", "21-24시"),
            ("sales_amt_age20", "20대 매출"), ("sales_amt_age30", "30대 매출")]
    tabs["상권 매출"] = [[h for _, h in cols]] + [
        [f"{r['quarter'][:4]}년 {r['quarter'][4:]}분기" if k == "quarter" else r[k] if k == "service_name" else _num(r[k])
         for k, _ in cols] for r in sorted(rows, key=lambda r: r["quarter"])]
    rows = list(csv.DictReader((DATA / "stores_sbiz.csv").open(encoding="utf-8")))
    cols = [("name", "상호"), ("branch_name", "지점명"), ("class_name", "소분류"), ("road_addr", "도로명주소"),
            ("floor", "층"), ("trdar_name", "상권"), ("role", "용도"), ("franchise_flag", "프랜차이즈(지점명 있음)"),
            ("snapshot", "기준일")]
    role = {"primary": "대학로 상권 (주 비교)", "comparison": "인근 상권 (비교)", "context": "주변 (참고)"}
    tabs["고깃집 목록"] = [[h for _, h in cols]] + [
        [role.get(r.get(k), r.get(k)) if k == "role" else ("예" if r.get(k) == "1" else "아니오") if k == "franchise_flag"
         else r.get(k, "") for k, _ in cols] for r in rows]
    d = json.loads((DATA / "survival_wide.json").read_text(encoding="utf-8"))
    head = ["집단", "표본 수", "폐업 수", "24개월 안 폐업 비율", "생존 중앙값(개월)", "면적 중앙값(평)",
            "12개월 생존율", "24개월 생존율", "36개월 생존율", "60개월 생존율", "메모"]
    labels = {"pork_name_only": "돼지고기 구이집 (상호로 판별)", "pork_like": "돼지고기 구이집 (상호+상가정보 매칭)",
              "hansik_all": "한식 전체", "hansik_not_pork": "돼지고기 구이집 뺀 한식", "all_uptae": "모든 업태"}
    notes = {"pork_name_only": "주 비교 대상, 상호 기준이라 편향 적고 잡음 많음", "pork_like": "지금 영업 중인 점포만 매칭되어 생존 쪽으로 편향"}
    srows = [head]
    for k, lab in labels.items():
        v = d.get(k) or {}
        tab = {t.get("month"): t.get("survival") for t in v.get("table") or []}

        def at(mo):
            ks = [m for m in tab if m is not None and m <= mo]
            return round(tab[max(ks)], 3) if ks else None
        srows.append([lab, v.get("n"), v.get("events"), v.get("share_closed_within_24m"), v.get("median_months"),
                      v.get("median_area_pyeong"), at(12), at(24), at(36), at(60), notes.get(k, "")])
    lr = (d.get("logrank") or {}).get("pork_name_only_vs_other_hansik") or {}
    srows += [[], [f"범위: {', '.join(d.get('areas') or [])}, {d.get('cohort', '')}"],
              [f"로그순위 검정 (상호 기준 돼지고기 구이집 vs 나머지 한식) p값 {lr.get('p_value', lr.get('p', '-'))}"]]
    tabs["생존 요약"] = srows
    src = _sources()
    tabs["출처"] = [["출처ID", "제목", "발행처/저자", "연도", "URL", "유형", "신뢰도 메모"]] + \
        [[s["source_id"], s["title"], s["publisher_author"], s["year"], s["url"], s["type"], s["reliability_note"]] for s in src]
    return tabs


def _rgb(h: str) -> dict:
    return {"red": int(h[0:2], 16) / 255, "green": int(h[2:4], 16) / 255, "blue": int(h[4:6], 16) / 255}


def _width(sheet: int, i: int, px: int) -> dict:
    return {"updateDimensionProperties": {"range": {"sheetId": sheet, "dimension": "COLUMNS", "startIndex": i, "endIndex": i + 1},
                                          "properties": {"pixelSize": px}, "fields": "pixelSize"}}


def _fmt(sheet: int, r0: int, r1: int, c0: int, c1: int, fmt: dict, fields: str) -> dict:
    return {"repeatCell": {"range": {"sheetId": sheet, "startRowIndex": r0, "endRowIndex": r1, "startColumnIndex": c0,
                                     "endColumnIndex": c1}, "cell": {"userEnteredFormat": fmt}, "fields": fields}}


def format_requests(ids: dict[str, int], sizes: dict[str, tuple[int, int]]) -> list[dict]:
    reqs = []
    black, gray, white = _rgb("111111"), _rgb("F3F4F6"), _rgb("FFFFFF")
    for title, sid in ids.items():
        nrows, ncols = sizes[title]
        reqs.append(_fmt(sid, 0, max(nrows, 1) + 200, 0, max(ncols, 1),
                         {"textFormat": {"fontFamily": FONT, "fontSize": 10}, "verticalAlignment": "TOP"},
                         "userEnteredFormat(textFormat,verticalAlignment)"))
        if title == "안내":
            reqs.append(_fmt(sid, 0, 1, 0, 1, {"textFormat": {"fontFamily": FONT, "fontSize": 16, "bold": True}},
                             "userEnteredFormat.textFormat"))
            for r, row in enumerate(GUIDE):
                if row[0] in ("내 탭", "3단계: 찾기, 확인, 기입", "맞음과 수정", "반영 규칙", "지킬 것", "참고 탭 (읽기 전용)"):
                    reqs.append(_fmt(sid, r, r + 1, 0, 1, {"backgroundColor": black,
                                                          "textFormat": {"fontFamily": FONT, "bold": True, "foregroundColor": white}},
                                     "userEnteredFormat(backgroundColor,textFormat)"))
            reqs.append(_width(sid, 0, 900))
            reqs.append(_fmt(sid, 0, nrows, 0, 1, {"wrapStrategy": "WRAP"}, "userEnteredFormat.wrapStrategy"))
            continue
        reqs.append({"updateSheetProperties": {"properties": {"sheetId": sid, "gridProperties": {"frozenRowCount": 1}},
                                               "fields": "gridProperties.frozenRowCount"}})
        reqs.append(_fmt(sid, 0, 1, 0, ncols, {"backgroundColor": black, "wrapStrategy": "WRAP",
                                               "textFormat": {"fontFamily": FONT, "bold": True, "foregroundColor": white}},
                         "userEnteredFormat(backgroundColor,textFormat,wrapStrategy)"))
        if title in MEMBERS:
            reqs.append(_fmt(sid, 1, nrows, 0, 8, {"backgroundColor": gray}, "userEnteredFormat.backgroundColor"))
            reqs.append(_fmt(sid, 1, nrows, 15, 16, {"backgroundColor": gray}, "userEnteredFormat.backgroundColor"))
            reqs.append(_fmt(sid, 1, nrows, 8, 15, {"backgroundColor": white}, "userEnteredFormat.backgroundColor"))
            reqs.append(_fmt(sid, 1, nrows, 1, 2, {"wrapStrategy": "WRAP"}, "userEnteredFormat.wrapStrategy"))
            reqs.append(_fmt(sid, 1, nrows, 7, 8, {"wrapStrategy": "WRAP"}, "userEnteredFormat.wrapStrategy"))
            reqs.append(_fmt(sid, 1, nrows, 13, 14, {"wrapStrategy": "WRAP"}, "userEnteredFormat.wrapStrategy"))
            reqs.append({"setDataValidation": {
                "range": {"sheetId": sid, "startRowIndex": 1, "endRowIndex": nrows, "startColumnIndex": 8, "endColumnIndex": 9},
                "rule": {"condition": {"type": "ONE_OF_LIST", "values": [{"userEnteredValue": v} for v in CHOICES]},
                         "showCustomUi": True, "strict": True}}})
            bad = ('=AND($I2="수정", OR(NOT(ISNUMBER($K2)), AND(ISNUMBER($J2), $J2>$K2), AND(ISNUMBER($L2), $K2>$L2), '
                   'AND($J2<>"", NOT(ISNUMBER($J2))), AND($L2<>"", NOT(ISNUMBER($L2)))))')
            reqs.append({"addConditionalFormatRule": {"index": 0, "rule": {
                "ranges": [{"sheetId": sid, "startRowIndex": 1, "endRowIndex": nrows, "startColumnIndex": 9, "endColumnIndex": 12}],
                "booleanRule": {"condition": {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": bad}]},
                                "format": {"backgroundColor": _rgb("FDE2E1")}}}}})
            for i, px in enumerate(WIDTHS):
                reqs.append(_width(sid, i, px))
            reqs.append({"addProtectedRange": {"protectedRange": {
                "range": {"sheetId": sid, "startRowIndex": 0, "endRowIndex": nrows, "startColumnIndex": 0, "endColumnIndex": 8},
                "description": "현재 값은 저장소에서 갱신됨 (--refresh)", "warningOnly": True}}})
            reqs.append({"addProtectedRange": {"protectedRange": {
                "range": {"sheetId": sid, "startRowIndex": 0, "endRowIndex": 1, "startColumnIndex": 8, "endColumnIndex": 16},
                "description": "머리글", "warningOnly": True}}})
        else:
            for i in range(ncols):
                reqs.append(_width(sid, i, 260 if title == "출처" and i in (1, 4, 6) else 140))
            reqs.append({"addProtectedRange": {"protectedRange": {"range": {"sheetId": sid},
                                                                  "description": "참고용 (읽기 전용)", "warningOnly": True}}})
    return reqs


def create() -> dict:
    items, cur = member_items(), current_values()
    content = {"안내": GUIDE}
    for m in MEMBERS:
        content[m] = [HEAD] + member_rows(m, items, cur)
    content.update(reference_tabs())
    res = gws("sheets", "spreadsheets", "create", body={
        "properties": {"title": TITLE, "locale": "ko_KR"},
        "sheets": [{"properties": {"title": t, "index": i}} for i, t in enumerate(content)]})
    sid = res["spreadsheetId"]
    ids = {s["properties"]["title"]: s["properties"]["sheetId"] for s in res["sheets"]}
    gws("sheets", "spreadsheets", "values", "batchUpdate", params={"spreadsheetId": sid},
        body={"valueInputOption": "RAW", "data": [{"range": f"'{t}'!A1", "values": rows} for t, rows in content.items()]})
    sizes = {t: (len(rows), max((len(r) for r in rows), default=1)) for t, rows in content.items()}
    gws("sheets", "spreadsheets", "batchUpdate", params={"spreadsheetId": sid},
        body={"requests": format_requests(ids, sizes)})
    meta = gws("drive", "files", "get", params={"fileId": sid, "fields": "parents"})
    gws("drive", "files", "update", params={"fileId": sid, "addParents": FOLDER_ID,
                                            "removeParents": ",".join(meta.get("parents") or []), "fields": "id,parents"})
    rec = {"spreadsheet_id": sid, "url": f"https://docs.google.com/spreadsheets/d/{sid}/edit", "title": TITLE,
           "folder_id": FOLDER_ID, "created": dt.date.today().isoformat(), "tabs": list(content),
           "rows": {m: len(content[m]) - 1 for m in MEMBERS}}
    SHEET_JSON.parent.mkdir(parents=True, exist_ok=True)
    SHEET_JSON.write_text(json.dumps(rec, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return rec


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--create", action="store_true", help="create the sheet in the Team 1 folder (once)")
    g.add_argument("--refresh", action="store_true", help="rewrite the 현재 columns from the workbook")
    g.add_argument("--apply", action="store_true", help="write the patch and rebuild the workbook")
    g.add_argument("--dry-run", action="store_true", help="the default: print what would change, write nothing")
    ap.add_argument("--stamp", action="store_true", help="with --apply: write 반영됨 dates on applied rows")
    a = ap.parse_args(argv)
    if a.create:
        if SHEET_JSON.exists():
            sys.exit(f"already created: {json.loads(SHEET_JSON.read_text(encoding='utf-8'))['url']}")
        rec = create()
        print(f"created {rec['url']} rows {rec['rows']}")
        return
    sid = sheet_id()
    if a.refresh:
        print(f"refreshed {refresh(sid)} rows (현재 columns only)")
        return
    result = collect(read_member_tabs(sid), current_values(), _sources())
    print(report(result))
    if a.apply:
        apply(result)
        if a.stamp:
            stamp(sid, result["applied"])
    else:
        print("dry run: nothing written (use --apply)")


if __name__ == "__main__":
    main()
