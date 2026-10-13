"""--apply-patch and --out for research/build_workbook.py (runs the script; writes only to tmp_path)."""
import hashlib
import subprocess
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / "research" / "build_workbook.py"
HERO_FILES = [ROOT / "research" / f for f in ("guesstimate.xlsx", "hero_base.yaml", "shocks.yaml")]


def digest(paths):
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}


def run(*args):
    return subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], cwd=ROOT, capture_output=True, text=True)


def test_out_writes_twin_without_touching_hero(tmp_path):
    before = digest(HERO_FILES)
    r = run("--out", tmp_path / "T99")
    assert r.returncode == 0, r.stderr
    for f in ("guesstimate.xlsx", "hero_base.yaml", "shocks.yaml"):
        assert (tmp_path / "T99" / f).exists()
    assert digest(HERO_FILES) == before


def test_helper_patch_applies_reviewed_rows_and_refuses_unreviewed(tmp_path):
    patch = tmp_path / "accepted_patch.yaml"
    patch.write_text(yaml.safe_dump({
        "_meta": {"generated_by": "test", "rows": 2},
        "cst_banchan_cost_per_cover": {"low": 900, "base": 1500, "high": 2100, "reasoning_chain": "테스트 검토값",
                                       "source_id": None, "confidence": "L", "filled_by": "AI", "reviewed_by": "검토자A"},
        "cst_utilities": {"low": 1, "base": 2, "high": 3, "reasoning_chain": "검토 안 됨", "source_id": None,
                          "confidence": "L", "filled_by": "AI", "reviewed_by": ""},
        "cst_side_cost_ratio": {"low": 0.3, "base": 0.9, "high": 0.4, "reasoning_chain": "범위 오류", "source_id": None,
                                "confidence": "L", "filled_by": "AI", "reviewed_by": "검토자A"},
    }, allow_unicode=True), encoding="utf-8")
    out = tmp_path / "patched"
    r = run("--apply-patch", patch, "--out", out)
    assert r.returncode == 0, r.stderr
    assert "REFUSED cst_utilities: no reviewed_by" in r.stdout
    assert "REFUSED cst_side_cost_ratio: base outside [low, high]" in r.stdout
    hero = yaml.safe_load((out / "hero_base.yaml").read_text(encoding="utf-8"))
    assert hero["cst_banchan_cost_per_cover"]["base"] == 1500
    assert hero["cst_utilities"]["base"] == 1700000        # refused row unchanged
    assert hero["cst_side_cost_ratio"]["base"] == 0.35
    wb = load_workbook(out / "guesstimate.xlsx")
    log = [r for r in wb["ChangeLog"].iter_rows(min_row=2, values_only=True) if r[0]]
    assert {r[0] for r in log} == {"cst_banchan_cost_per_cover"}
    base_row = [r for r in log if r[1] == "base"][0]
    assert (base_row[2], base_row[3], base_row[5]) == (1200, 1500, "검토자A")
    hs = wb["HeroStore"]
    head = [c.value for c in hs[1]]
    row = next(r for r in hs.iter_rows(min_row=2, values_only=True) if r[0] == "cst_banchan_cost_per_cover")
    assert row[head.index("reviewed_by")] == "검토자A"


def test_interview_patch_uses_meta_reviewer_and_refuses_unknown_source(tmp_path):
    patch = tmp_path / "T01.yaml"
    patch.write_text(yaml.safe_dump({
        "_meta": {"twin_id": "T01", "interview_date": "2026-10-05", "interviewer": "멤버1", "reviewed_by": "멤버2",
                  "source_id": "S36"},
        "hero": {"drv_dwell_time": {"low": 80, "base": 100, "high": 130, "confidence": "M",
                                    "reasoning_chain": "T01 Q8", "source_id": "S36"},
                 "mkt_patience": {"low": 10, "base": 20, "high": 30, "confidence": "M", "reasoning_chain": "T01 Q9",
                                  "source_id": "INT-T01"}},
        "shocks": {"SH09": {"parameter": "drv_arrival_rate", "delta_low": -0.1, "delta_base": -0.15, "delta_high": -0.25,
                            "rationale": "T01 Q10", "source_id": "S12", "confidence": "M"}},
        "owner_policy": {"SH03": {"action": "reprice", "args": {"pct": 0.07}}},
    }, allow_unicode=True), encoding="utf-8")
    out = tmp_path / "T01"
    r = run("--apply-patch", patch, "--out", out)
    assert r.returncode == 0, r.stderr
    assert "REFUSED mkt_patience: unknown source_id ['INT-T01']" in r.stdout
    assert "not applied here (not workbook values): owner_policy" in r.stdout
    hero = yaml.safe_load((out / "hero_base.yaml").read_text(encoding="utf-8"))
    assert hero["drv_dwell_time"]["base"] == 100 and hero["mkt_patience"]["base"] == 25
    shocks = yaml.safe_load((out / "shocks.yaml").read_text(encoding="utf-8"))
    assert shocks["SH09"]["params"][0]["base"] == -0.15
    log = list(load_workbook(out / "guesstimate.xlsx")["ChangeLog"].iter_rows(min_row=2, values_only=True))
    assert any(r[0] == "SH09.drv_arrival_rate" and r[1] == "delta_base" and r[5] == "멤버2" and r[6] == "2026-10-05" for r in log)
