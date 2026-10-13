"""Load ontology, research inputs and processed public data into the DB. Idempotent (merge = upsert).

  python -m db.load                 -> DATABASE_URL (default sqlite:///data/twin.db)
  python -m db.load --selftest      -> a temp SQLite file, loaded twice, counts asserted
"""
from __future__ import annotations

import argparse
import json
import math
import sys
import tempfile
from datetime import date
from pathlib import Path

import pandas as pd
import yaml
from sqlalchemy import func, select

from . import models as M
from .session import REPO, get_engine, init_db, session_factory

ONTOLOGY = REPO / "ontology"
RESEARCH = REPO / "research"
PROCESSED = REPO / "data" / "processed"
INDUSTRY_YAML = REPO / "industry" / "pork_bbq.yaml"


def log(msg):
    print(msg, file=sys.stderr, flush=True)


def _clean(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    if isinstance(v, str) and v.strip() == "":
        return None
    return v


def _csv(path: Path) -> list[dict]:
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    return [{k: _clean(v) for k, v in r.items()} for r in df.to_dict("records")]


def _num(v):
    try:
        f = float(v)
        return None if math.isnan(f) else f
    except (TypeError, ValueError):
        return None


def _date(v):
    v = _clean(v)
    return date.fromisoformat(str(v)[:10]) if v else None


def _yn(v):
    v = _clean(v)
    return None if v is None else str(v).strip().upper() in ("Y", "1", "TRUE", "YES")


# --- loaders ---

def _benchmarks() -> list[dict]:
    """Benchmarks sheet of research/guesstimate.xlsx (read only), e.g. the 혜화동 vs 전국 rent comparison."""
    path = RESEARCH / "guesstimate.xlsx"
    if not path.exists():
        return []
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Benchmarks" not in wb.sheetnames:
        return []
    rows = list(wb["Benchmarks"].iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    return [{k: (v if not isinstance(v, float) or v == v else None) for k, v in zip(head, r)}
            for r in rows[1:] if r and r[0]]


def load_industry(s, cfg, profile):
    s.merge(M.Industry(industry_id=cfg["industry"], label_ko=cfg.get("label_ko"), config=cfg,
                       meta={"profile_meta": profile.get("meta"), "survival": profile.get("survival"),
                             "survival_scope": profile.get("survival_scope", "narrow"),
                             "survival_wide": profile.get("survival_wide"),
                             "benchmarks": json.loads(json.dumps(_benchmarks(), ensure_ascii=False, default=str))}))
    return 1


def load_ontology(s, industry_id):
    n = {}
    rows = _csv(ONTOLOGY / "sources.csv")
    for r in rows:
        s.merge(M.Source(**r))
    n["source"] = len(rows)
    rows = _csv(ONTOLOGY / "criteria.csv")
    for r in rows:
        s.merge(M.Criterion(
            criterion_id=r["criterion_id"], industry_id=industry_id, name_ko=r["name_ko"], name_en=r["name_en"],
            category=r["category"], driver_node=r["driver_node"], unit=r["unit"], data_type=r["data_type"],
            method=r["method"], source_ids=r["source_id"], industry_scope=r["industry_scope"],
            in_sim=_yn(r["in_sim"]), notes=r["notes"]))
    n["criterion"] = len(rows)
    s.flush()
    rows = _csv(ONTOLOGY / "edges.csv")
    for r in rows:
        s.merge(M.CriterionEdge(from_id=r["from_id"], to_id=r["to_id"], sign=r["sign"],
                                mechanism_ko=r["mechanism_ko"], source_ids=r["source_id"],
                                confidence=r["confidence"]))
    n["criterion_edge"] = len(rows)
    return n


def _workbook_hero_meta() -> dict:
    """filled_by / reviewed_by / reasoning from research/guesstimate.xlsx (members edit reviewed_by there)."""
    path = RESEARCH / "guesstimate.xlsx"
    if not path.exists():
        return {}
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "HeroStore" not in wb.sheetnames:
        return {}
    rows = list(wb["HeroStore"].iter_rows(values_only=True))
    head = list(rows[0])
    return {r[head.index("criterion_id")]: dict(zip(head, r)) for r in rows[1:] if r and r[0]}


def load_hero(s, industry_id):
    hero = yaml.safe_load((RESEARCH / "hero_base.yaml").read_text(encoding="utf-8"))
    wb = _workbook_hero_meta()
    known = set(s.scalars(select(M.Criterion.criterion_id)))
    n, skipped = 0, []
    for cid, v in hero.items():
        if cid.startswith("_"):
            continue
        if cid not in known:
            skipped.append(cid)
            continue
        w = wb.get(cid, {})
        method = v.get("method")
        s.merge(M.HeroAttribute(
            industry_id=industry_id, criterion_id=cid,
            low=None if v.get("low") is None else str(v["low"]),
            base=None if v.get("base") is None else str(v["base"]),
            high=None if v.get("high") is None else str(v["high"]),
            low_num=_num(v.get("low")), base_num=_num(v.get("base")), high_num=_num(v.get("high")),
            unit=v.get("unit"), method=method, source_id=v.get("source_id"), confidence=v.get("confidence"),
            reasoning_chain=_clean(w.get("reasoning_chain")),
            filled_by=_clean(w.get("filled_by")) or ("formula" if method == "derived" else "AI"),
            reviewed_by=_clean(w.get("reviewed_by"))))
        n += 1
    if skipped:
        log(f"  WARNING: hero_base.yaml ids not in criteria.csv, skipped: {skipped}")
    s.merge(M.Scenario(scenario_id="hero_base", industry_id=industry_id, name="Hero store base case",
                       description=(hero.get("_meta") or {}).get("store"),
                       params={"meta": hero.get("_meta"), "pnl_base": hero.get("_pnl_base")}))
    return {"hero_attribute": n, "scenario": 1, "hero_skipped": len(skipped)}


def _shock_rows() -> tuple[list[dict], str]:
    y = RESEARCH / "shocks.yaml"
    if y.exists():
        data = yaml.safe_load(y.read_text(encoding="utf-8"))
        if isinstance(data, dict) and "shocks" in data:
            data = data["shocks"]
        if isinstance(data, dict):
            data = [{"shock_id": k, **(v or {})} for k, v in data.items() if not str(k).startswith("_")]
        return list(data or []), "research/shocks.yaml"
    path = RESEARCH / "guesstimate.xlsx"
    if path.exists():
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "Shocks" in wb.sheetnames:
            rows = list(wb["Shocks"].iter_rows(values_only=True))
            return [dict(zip(rows[0], r)) for r in rows[1:] if r and r[0]], "guesstimate.xlsx:Shocks"
    return [], "none"


def load_shocks(s, industry_id):
    rows, origin = _shock_rows()
    for r in rows:
        # shocks.yaml nests one or more parameter deltas under `params`; the flat columns hold the first
        # (the full list stays in `spec`), and `parameter` lists every parameter the shock touches.
        ps = r.get("params") if isinstance(r.get("params"), list) else []
        flat = {**(ps[0] if ps else {}), **{k: v for k, v in r.items() if k != "params"}}
        if ps:
            flat.setdefault("delta_low", ps[0].get("low"))
            flat.setdefault("delta_base", ps[0].get("base"))
            flat.setdefault("delta_high", ps[0].get("high"))
            flat["parameter"] = ";".join(str(x.get("parameter")) for x in ps)
        g = lambda *ks: next((flat[k] for k in ks if _clean(flat.get(k)) is not None), None)  # noqa: E731
        sid = str(g("shock_id", "id"))
        s.merge(M.Shock(
            shock_id=sid, industry_id=industry_id, driver=g("driver"), name_ko=g("name_ko", "name", "label_ko", "label"),
            parameter=g("parameter", "param"), delta_type=g("delta_type"),
            delta_low=_num(g("delta_low")), delta_base=_num(g("delta_base", "delta")), delta_high=_num(g("delta_high")),
            mechanism=g("mechanism", "mechanism_ko"), source_id=g("source_id"), confidence=g("confidence"),
            spec=json.loads(json.dumps(r, ensure_ascii=False, default=str))))
    return {"shock": len(rows), "shock_origin": origin}


def load_district(s, industry_id, profile):
    trdar = {r["trdar_code"]: r for r in _csv(PROCESSED / "seoul_trdar.csv")}
    geo = {}
    gpath = PROCESSED / "trdar_boundaries.geojson"
    if gpath.exists():
        for f in json.loads(gpath.read_text(encoding="utf-8"))["features"]:
            geo[f["properties"]["trdar_code"]] = f["geometry"]
    nq = 0
    for code, d in profile["trdar"].items():
        t = trdar.get(code, {})
        s.merge(M.District(
            trdar_code=code, industry_id=industry_id, trdar_name=d["trdar_name"], trdar_type=d["trdar_type"],
            role=d["role"], adstrd_name=d["adstrd_name"], signgu_name=d["signgu_name"], area_m2=_num(d["area_m2"]),
            centroid_lon=_num(t.get("centroid_lon")), centroid_lat=_num(t.get("centroid_lat")),
            geometry=geo.get(code),
            summary={k: d[k] for k in ("sbiz_stores", "licenses", "store_license_match", "annual")}))
        s.flush()
        for q, v in d["quarters"].items():
            shares = {k: v.get(k) for k in ("dow_share", "timeband_share", "age_share", "sex_share",
                                            "age_split_coverage", "flpop_age_share", "flpop_timeband_share",
                                            "facilities") if k in v}
            s.merge(M.DistrictQuarter(
                trdar_code=code, quarter=q, shares=shares,
                **{c: v.get(c) for c in (
                    "hansik_sales_krw", "hansik_sales_count", "avg_ticket_krw", "stores_total", "stores_independent",
                    "stores_franchise", "stores_opened", "stores_closed", "sales_per_store_month_krw",
                    "evening_share_17_24", "lunch_share_11_14", "weekend_share", "flpop_total", "workers_total",
                    "residents_total", "households_total")}))
            nq += 1
    return {"district": len(profile["trdar"]), "district_quarter": nq}


def load_stores(s):
    match = {r["store_id"]: r for r in _csv(PROCESSED / "store_license_match.csv")}
    rows = _csv(PROCESSED / "stores_sbiz.csv")
    for r in rows:
        m = match.get(r["store_id"], {})
        s.merge(M.Store(
            store_id=r["store_id"], trdar_code=r["trdar_code"], name=r["name"], branch_name=r["branch_name"],
            class_code=r["class_code"], class_name=r["class_name"], role=r["role"],
            franchise_flag=_yn(r["franchise_flag"]), road_addr=r["road_addr"], floor=r["floor"],
            lon=_num(r["lon"]), lat=_num(r["lat"]), snapshot=r["snapshot"],
            license_id=m.get("license_id"), match_confidence=m.get("match_confidence"),
            match_name_score=_num(m.get("name_score")), match_dist_m=_num(m.get("dist_m"))))
    return {"store": len(rows)}


def load_licenses(s):
    rows = _csv(PROCESSED / "licenses.csv")
    for r in rows:
        s.merge(M.StoreLicense(
            license_id=r["license_id"], trdar_code=r["trdar_code"], name=r["name"], uptae=r["uptae"],
            hansik=_yn(r["hansik"]), pork_name_hint=_yn(r["pork_name_hint"]), status=r["status"],
            open_date=_date(r["open_date"]), close_date=_date(r["close_date"]), censor_date=_date(r["censor_date"]),
            months_open=_num(r["months_open"]), event=int(float(r["event"])),
            closed_within_24m=None if _num(r.get("closed_within_24m")) is None else int(float(r["closed_within_24m"])),
            area_m2=_num(r["area_m2"]), area_pyeong=_num(r["area_pyeong"]), road_addr=r["road_addr"],
            x_5181=_num(r["x_5181"]), y_5181=_num(r["y_5181"]), coord_source=r["coord_source"]))
    return {"store_license": len(rows)}


def load_all(engine=None) -> dict:
    engine = init_db(engine or get_engine())
    cfg = yaml.safe_load(INDUSTRY_YAML.read_text(encoding="utf-8"))
    profile = json.loads((PROCESSED / "district_profile.json").read_text(encoding="utf-8"))
    ind = cfg["industry"]
    counts = {}
    with session_factory(engine)() as s, s.begin():
        counts["industry"] = load_industry(s, cfg, profile)
        s.flush()
        counts.update(load_ontology(s, ind))
        s.flush()
        counts.update(load_hero(s, ind))
        counts.update(load_shocks(s, ind))
        counts.update(load_district(s, ind, profile))
        s.flush()
        counts.update(load_stores(s))
        counts.update(load_licenses(s))
    return counts


def table_counts(engine) -> dict:
    with session_factory(engine)() as s:
        return {t.__tablename__: s.scalar(select(func.count()).select_from(t)) for t in (
            M.Industry, M.Source, M.Criterion, M.CriterionEdge, M.HeroAttribute, M.Scenario, M.Shock,
            M.District, M.DistrictQuarter, M.Store, M.StoreLicense)}


def selftest():
    with tempfile.TemporaryDirectory() as d:
        eng = get_engine(f"sqlite:///{d}/selftest.db")
        first = load_all(eng)
        c1 = table_counts(eng)
        load_all(eng)
        c2 = table_counts(eng)
        assert c1 == c2, f"not idempotent: {c1} vs {c2}"
        for t in ("source", "criterion", "criterion_edge", "store", "store_license", "district_quarter"):
            assert c1[t] == first[t], f"{t}: loaded {first[t]} rows, table has {c1[t]} (duplicate keys in input?)"
        assert c1["hero_attribute"] == first["hero_attribute"] > 0
        assert c1["district"] == len(yaml.safe_load(INDUSTRY_YAML.read_text())["trdar"]["primary"]) + \
            len(yaml.safe_load(INDUSTRY_YAML.read_text())["trdar"]["fringe"])
        with session_factory(eng)() as s:
            q = s.get(M.DistrictQuarter, ("3120012", "20254"))
            assert q is not None and q.stores_total == 191 and q.hansik_sales_krw > 0
            assert s.scalar(select(func.count()).select_from(M.Store).where(
                M.Store.trdar_code == "3120012", M.Store.class_code == "I20107")) == 27
            orphans = s.scalar(select(func.count()).select_from(M.CriterionEdge).where(
                ~M.CriterionEdge.to_id.in_(select(M.Criterion.criterion_id))))
            assert orphans == 0
        eng.dispose()
        log(f"DB LOAD SELFTEST PASS (loaded twice, identical counts): {c1} | shocks from {first['shock_origin']}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        selftest()
    else:
        counts = load_all()
        log(f"loaded: {counts}")
        log(f"tables: {table_counts(get_engine())}")
